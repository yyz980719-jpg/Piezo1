from __future__ import annotations

import subprocess
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(r"E:\codex\2026-08-31\yu")
PACKAGE = ROOT / "outputs" / "PIEZO1_OA_V10_V8_FinalExecution_Audit_20260831"
REPOSITORY = ROOT / "outputs" / "PIEZO1_OA_public_repository"
MAIN_IN = PACKAGE / "PIEZO1_OA_Main_Manuscript_V13_Corrected_SOP_Audit_HOLD.docx"
SUPP_IN = PACKAGE / "PIEZO1_OA_Supplementary_Materials_V11_Corrected_SOP_Audit_HOLD.docx"
STROBE_IN = PACKAGE / "PIEZO1_OA_STROBE-MR_Checklist_Completed_V2.docx"
MAIN_OUT = PACKAGE / "PIEZO1_OA_Main_Manuscript_V14_Final_SOP_Review.docx"
SUPP_OUT = PACKAGE / "PIEZO1_OA_Supplementary_Materials_V12_Final_SOP_Review.docx"
STROBE_OUT = PACKAGE / "PIEZO1_OA_STROBE-MR_Checklist_Completed_V3.docx"
GIT = Path(
    r"C:\Users\Administrator\.cache\codex-runtimes\codex-primary-runtime"
    r"\dependencies\native\git\cmd\git.exe"
)


def set_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell(cell, text: str) -> None:
    set_paragraph(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph(paragraph, "")


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise KeyError(f"Paragraph not found: {prefix}")


def replace_document_text(document: Document, replacements: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        if text != paragraph.text:
            set_paragraph(paragraph, text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                for old, new in replacements.items():
                    text = text.replace(old, new)
                if text != cell.text:
                    set_cell(cell, text)


def clean_properties(document: Document, title: str, subject: str) -> None:
    props = document.core_properties
    props.title = title
    props.subject = subject
    props.author = "Yang Yunze"
    props.last_modified_by = "Yang Yunze"
    props.comments = ""
    props.revision = 1
    now = datetime.now(timezone.utc)
    props.created = now
    props.modified = now


def repository_commit() -> str:
    result = subprocess.run(
        [str(GIT), "rev-parse", "HEAD"],
        cwd=REPOSITORY,
        check=True,
        capture_output=True,
        text=True,
    )
    return result.stdout.strip()


def revise_main(commit: str) -> Document:
    doc = Document(MAIN_IN)
    replacements = {
        "independent gene-set specifications": "unclosed gene-set specifications",
        "independent gene-set score": "unclosed gene-set score",
        "independent gene-set treatment": "unclosed gene-set treatment",
        "final-audit recall-gated sensitivity": "recall-gated sensitivity",
        "final-audit sensitivity analysis": "final sensitivity analysis",
    }
    replace_document_text(doc, replacements)

    set_paragraph(
        find_paragraph(doc, "GSE104782 cells were assigned"),
        "GSE104782 cells were assigned to seven published chondrocyte states. The continuous axis was defined as z(mean hypertrophic markers) minus z(mean homeostatic markers), using COL10A1, RUNX2, IBSP, ALPL, MMP13, SPP1 and POSTN for the hypertrophic set and COL2A1, SOX9, ACAN, CHAD and HAPLN1 for the homeostatic set. CRTL1 is a historical alias of HAPLN1 and was not treated as an additional independent marker. PIEZO1 was not included in either set. PIEZO1-axis association was estimated within donor and summarized across ten donors, and Q1/Q4 groups were defined within each donor. Cell-level tests were descriptive; donor-blocked pseudobulk models were used for categorical contrasts with transcriptome-wide correction. The paired Q1-Q4 Wilcoxon signed-rank sensitivity test used the asymptotic normal approximation with continuity correction.",
    )
    set_paragraph(
        find_paragraph(doc, "State programmes were scored"),
        "State programmes were scored by one final standardized pipeline distinct from the disease-level DEG preprocessing. GSE104782-derived top-100 state markers were fitted after probe-mean aggregation, conditional conversion of log2 bulk expression to the linear scale and gene-wise standardization; no additional quantile normalization was applied. Non-negative L-BFGS-B weights were then closed to sum one. Table S6 reports the raw closed-weight layer, while the zero-handling, centred-log-ratio, additive-log-ratio, pairwise-log-ratio and unclosed gene-set specifications in Table S13 use the same marker construction and authoritative programme preprocessing. The homeostatic-minus-pre-hypertrophic contrast, rather than any isolated coordinate, was the principal cross-dataset comparison.",
    )
    set_paragraph(
        find_paragraph(doc, "PIEZO1 was detected in 59.6%"),
        "PIEZO1 was detected in 59.6% of the 1,464 GSE104782 chondrocytes. Descriptively, mean expression was greatest in regulatory and homeostatic cells, but categorical state differences were sensitive to the level of analysis. The cell-level seven-state Kruskal-Wallis test was significant (P=3.9e-6), whereas the donor-level Friedman analysis of states represented in at least eight donors was not (chi-square=1.36, df=5, P=0.929; Kendall W=0.034), and the donor-level HomC-versus-HTC paired difference was not significant (mean 0.042, P=0.49). The donor-blocked pseudobulk subtype omnibus test was F(6, 38.81)=3.252, P=0.0110 and FDR=0.0663, and therefore did not pass transcriptome-wide correction. One pairwise contrast, RegC versus FC, remained significant after correction (log2 fold change 1.799, 95% CI 0.741 to 2.856, FDR=0.029); because the omnibus test was not significant after correction, it is reported as exploratory and supports no main conclusion.",
    )
    set_paragraph(
        find_paragraph(doc, "3.5 PIEZO1 direction is context dependent"),
        "3.5 PIEZO1 direction is context dependent, with donor-heterogeneous within-state and compositional contributions",
    )
    set_paragraph(
        find_paragraph(doc, "Code availability."),
        "Code availability. A versioned local Git repository contains the analytical scripts, software versions, derived results and data-to-result manifest used for this revision. Local release v1.2.0 is pinned to commit "
        + commit
        + ". The repository and DOI-bearing archive are not yet publicly available.",
    )
    clean_properties(
        doc,
        "PIEZO1 OA Main Manuscript V14",
        "Final-SOP reviewed manuscript; public repository and DOI remain unavailable",
    )
    return doc


def revise_supplement() -> Document:
    doc = Document(SUPP_IN)
    replacements = {
        "independent gene-set score": "unclosed gene-set score",
        "independent gene-set sensitivity": "unclosed gene-set sensitivity",
        "independent score": "unclosed score",
        "L4 independent gene-set score": "L4 unclosed gene-set score",
        "L4_independent_gs_score": "L4_unclosed_gs_score",
        "final-audit sensitivity analysis": "final sensitivity analysis",
        "final-audit sensitivity": "final sensitivity",
        "final-audit recall-gated sensitivity": "recall-gated sensitivity",
    }
    replace_document_text(doc, replacements)

    set_paragraph(
        find_paragraph(doc, "Continuous axis evidence is primary"),
        "Continuous axis evidence is primary because it supplies one correlation and one quartile contrast per donor. The axis was z(mean hypertrophic markers) minus z(mean homeostatic markers). Hypertrophic markers were COL10A1, RUNX2, IBSP, ALPL, MMP13, SPP1 and POSTN; homeostatic markers were COL2A1, SOX9, ACAN, CHAD and HAPLN1. CRTL1 is a historical alias of HAPLN1 and was not treated as an additional independent marker. PIEZO1 was absent from both sets. Q1 and Q4 were defined within donor. Pseudobulk donor-by-state analysis is supportive and requires at least five cells per donor-state combination. Cell-level P values are descriptive only. No claim of a stable ordering across all seven states is made because the donor-level Friedman test was null. The paired Q1-Q4 Wilcoxon signed-rank sensitivity used the asymptotic normal approximation with continuity correction.",
    )
    set_paragraph(
        find_paragraph(doc, "Deposited matrices were inspected for scale"),
        "The preprocessing described in S4 applies to disease-level and differential-expression analyses. Deposited series matrices were inspected for scale; matrices exceeding the raw-intensity threshold were log2 transformed, while already log2-scale matrices were retained. The authoritative GSE51588 DEG script did not apply additional quantile normalization. Probes were mapped to gene symbols using the platform annotation supplied with each series. Where several probes mapped to one gene, the probe with the highest mean expression was retained for the DEG pipeline. This S4 rule does not define programme scoring, which follows the separate authoritative pipeline in S5. No cross-platform batch correction was applied because inference remained within dataset.",
    )
    set_paragraph(
        find_paragraph(doc, "For each GSE104782 state"),
        "For each GSE104782 state, fold changes were computed against the mean of the remaining states and the top 100 genes were retained after requiring mean normalized expression above 0.05, giving 690 unique signature genes. The authoritative final programme pipeline converted the GSE104782 normalized matrix back to the linear scale for centroid construction. Bulk probes were aggregated to gene symbols by their mean, converted from log2 to linear scale when appropriate, and standardized gene-wise across samples; no additional quantile normalization was applied. Non-negative L-BFGS-B weights were fitted and then closed to sum one. Table S6 reports this L0 raw-weight layer, and Table S13 applies all zero-handling, log-ratio and unclosed gene-set sensitivity layers to the same final marker construction and programme preprocessing. For GSE152805, query cells and reference centroids were standardized over frozen genes and assigned by the highest unconstrained similarity projection score; these scores are arbitrary units, not proportions. PIEZO1 was absent from all top-100 marker sets.",
    )

    # Supplementary Table S5: moderated omnibus F degrees of freedom from the
    # unchanged authoritative limma model object.
    table_s5 = doc.tables[6]
    for row in table_s5.rows[1:]:
        if row.cells[0].text == "Pseudobulk subtype omnibus":
            set_cell(row.cells[2], "F(6, 38.81)=3.252")
            set_cell(row.cells[3], "0.0110; FDR 0.0663")

    # Supplementary Table S17: current transformation and outputs.
    table_s17 = doc.tables[19]
    for row in table_s17.rows[1:]:
        if row.cells[0].text == "Final execution document transformation":
            set_cell(row.cells[1], "scripts/89_apply_v13_v11_final_sop.py")
            set_cell(
                row.cells[2],
                "PIEZO1_OA_Main_Manuscript_V14_Final_SOP_Review.docx; "
                "PIEZO1_OA_Supplementary_Materials_V12_Final_SOP_Review.docx; "
                "PIEZO1_OA_STROBE-MR_Checklist_Completed_V3.docx",
            )

    # Supplementary Table S20: canonicalize the historical alias without
    # changing the five-gene homeostatic axis.
    table_s20 = doc.tables[22]
    set_cell(table_s20.rows[1].cells[0], "homeostatic_markers_requested")
    set_cell(table_s20.rows[1].cells[1], "COL2A1; SOX9; ACAN; CHAD; HAPLN1")
    set_cell(table_s20.rows[1].cells[2], "canonical HGNC/NCBI symbols")
    set_cell(table_s20.rows[2].cells[0], "homeostatic_markers")
    set_cell(table_s20.rows[2].cells[1], "COL2A1; SOX9; ACAN; CHAD; HAPLN1")
    set_cell(table_s20.rows[2].cells[2], "higher values decrease the axis score")
    set_cell(table_s20.rows[3].cells[0], "homeostatic_alias_resolution")
    set_cell(table_s20.rows[3].cells[1], "CRTL1 -> HAPLN1")
    set_cell(
        table_s20.rows[3].cells[2],
        "CRTL1 is a historical alias of HAPLN1 and was not counted as an additional marker",
    )

    # Keep the S16 title and header with the software table instead of leaving
    # a stranded heading/header at the foot of the preceding page.
    find_paragraph(doc, "Supplementary Table S16.").paragraph_format.page_break_before = True

    clean_properties(
        doc,
        "PIEZO1 OA Supplementary Materials V12",
        "Final-SOP reviewed supplement; public repository and DOI remain unavailable",
    )
    return doc


def revise_strobe() -> Document:
    doc = Document(STROBE_IN)
    table = doc.tables[0]
    set_cell(
        table.rows[54].cells[4],
        "Public source accessions, derived-result map, scripts, exact software versions and local release v1.2.0 are reported. A public repository URL and DOI-bearing archive are not yet available.",
    )
    clean_properties(
        doc,
        "Completed STROBE-MR checklist - PIEZO1 OA (V3)",
        "Updated checklist accompanying the final-SOP reviewed manuscript",
    )
    return doc


def main() -> None:
    commit = repository_commit()
    main_doc = revise_main(commit)
    supp_doc = revise_supplement()
    strobe_doc = revise_strobe()
    main_doc.save(MAIN_OUT)
    supp_doc.save(SUPP_OUT)
    strobe_doc.save(STROBE_OUT)
    print(MAIN_OUT)
    print(SUPP_OUT)
    print(STROBE_OUT)


if __name__ == "__main__":
    main()
