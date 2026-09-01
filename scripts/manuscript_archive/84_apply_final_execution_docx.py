from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"E:\codex\2026-08-31\yu")
SOURCE_MAIN = ROOT / "outputs" / "PIEZO1_OA_Main_Manuscript_V10_Final_QA.docx"
SOURCE_SUPP = ROOT / "outputs" / "PIEZO1_OA_Supplementary_Materials_V8_Final_QA.docx"
PACKAGE = ROOT / "outputs" / "PIEZO1_OA_V10_V8_FinalExecution_Audit_20260831"
MAIN_OUT = PACKAGE / "PIEZO1_OA_Main_Manuscript_V13_Corrected_SOP_Audit_HOLD.docx"
SUPP_OUT = PACKAGE / "PIEZO1_OA_Supplementary_Materials_V11_Corrected_SOP_Audit_HOLD.docx"
STROBE_TEMPLATE = ROOT / "work" / "STROBE-MR_official_fillable.docx"
STROBE_OUT = PACKAGE / "PIEZO1_OA_STROBE-MR_Checklist_Completed_V2.docx"
FIG1 = PACKAGE / "Figures" / "Figure1_Final_P0500.png"
FIGS1 = PACKAGE / "Figures" / "Figure_S1_Disease_Context_P0500.png"
FIGS15 = PACKAGE / "Figures" / "Figure_S15_Nested_StateTransfer_Validation.png"
FIG4 = PACKAGE / "Figures" / "Figure4_programme_robustness_direction_v11.png"
FIG5 = PACKAGE / "Figures" / "Figure5_MR_coloc_power_v11.png"
AUTHOR = "Yang Yunze"
ORCID = "0009-0003-5907-7037"
EMAIL = "631652214@qq.com"
AFFILIATION = ("Orthopedics and Traumatology, The First Affiliated Hospital of Heilongjiang "
               "University of Chinese Medicine, Heilongjiang University of Chinese Medicine, "
               "Harbin, Heilongjiang 150060, China")


def set_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def find_paragraph(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(prefix):
            return paragraph
    raise KeyError(f"Paragraph not found: {prefix}")


def set_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph(extra, "")


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def append_styled_row(table, values: list[str]) -> None:
    row = table.add_row()
    template = table.rows[-2]
    for idx, (cell, value) in enumerate(zip(row.cells, values)):
        current_tc_pr = cell._tc.tcPr
        if current_tc_pr is not None:
            cell._tc.remove(current_tc_pr)
        source_tc_pr = template.cells[idx]._tc.tcPr
        if source_tc_pr is not None:
            cell._tc.insert(0, deepcopy(source_tc_pr))
        set_cell(cell, value)


def replace_image_by_order(document: Document, image_index: int, image_path: Path) -> None:
    image_parts = []
    for paragraph in document.paragraphs:
        for blip in paragraph._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            part = document.part.rels[rid].target_part
            image_parts.append(part)
    image_parts[image_index]._blob = image_path.read_bytes()


def count_words(text: str) -> int:
    return len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b", text, flags=re.UNICODE))


def clean_properties(document: Document, title: str, subject: str) -> None:
    props = document.core_properties
    props.title = title
    props.subject = subject
    props.keywords = "PIEZO1; osteoarthritis; chondrocyte state; compositional data; single-cell RNA sequencing; Mendelian randomization"
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.comments = ""
    props.category = "Original research article"
    props.revision = 1
    now = datetime.now(timezone.utc)
    props.created = now
    props.modified = now


def revise_main() -> Document:
    doc = Document(SOURCE_MAIN)

    set_paragraph(find_paragraph(doc, "[Author names and degrees"),
        f"{AUTHOR} (ORCID: {ORCID})")
    set_paragraph(find_paragraph(doc, "[Affiliations"), AFFILIATION)
    set_paragraph(find_paragraph(doc, "Corresponding author:"),
        f"Corresponding author: {AUTHOR}, {AFFILIATION}. E-mail: {EMAIL}. ORCID: https://orcid.org/{ORCID}.")

    set_paragraph(find_paragraph(doc, "Objective."),
        "Objective. PIEZO1 is a mechanically gated cation channel implicated in cartilage mechanotransduction, but its reported expression in human osteoarthritis (OA) is inconsistent. We tested whether within-joint context is more strongly associated with PIEZO1 expression than case-control status, whether a chondrocyte-programme shift shows cross-dataset concordance under compositional testing, and whether genetically predicted whole-blood PIEZO1 expression influences OA risk under instrumental-variable assumptions.")
    set_paragraph(find_paragraph(doc, "Design."),
        "Design. Seven public transcriptomic datasets were integrated: paired subchondral bone, two synovial cohorts, an inflammation-gradient model, donor-annotated single-cell OA cartilage, paired preserved and affected cartilage, and paired lateral and medial cartilage single-cell profiles. Analyses used donor-paired differential expression, donor-level single-cell inference, programme scoring with compositional sensitivity analyses, reference-internal fully nested validation, external projection support, Mendelian randomization, colocalization and power analyses.")
    set_paragraph(find_paragraph(doc, "Results."),
        "Results. PIEZO1 showed no consistent OA-versus-control main effect. In OA subchondral bone, expression was lower medially (log2 fold change -0.504, P=0.000273, FDR=0.00189), but the disease-by-compartment interaction was not significant (P=0.147; covariate-adjusted P=0.81). Across 10 GSE104782 donors, PIEZO1 correlated weakly and negatively with the hypertrophic-minus-homeostatic axis (mean rho=-0.137, P=0.0035), and the donor-within Q1-Q4 difference was 0.123 (Wilcoxon P=0.0323). The homeostatic-versus-pre-hypertrophic contrast was concordant in 33 paired cartilage donors; GSE152805 supplied three-donor external projection support. Reference-internal fully nested validation exceeded the complete-pipeline permutation null (balanced accuracy 0.595; null 95th percentile 0.172; P=0.001). After excluding HTC, RegC and EC for recall below 0.60, the total medial-higher difference persisted in all three donors, but within-state expression and composition components were concordant in only two. Five whole-blood cis-eQTL instruments gave null primary estimates; conservative 80%-power minimum detectable odds ratios were 1.14, 1.14 and 1.09.")
    set_paragraph(find_paragraph(doc, "Conclusions."),
        "Conclusions. Human OA shows a consistent chondrocyte-programme balance shift across the analysed paired datasets, whereas PIEZO1 transcription is context dependent. The compartment difference has donor-heterogeneous within-state expression and composition contributions. Within the whole-blood instrument model, genetic estimates did not support a moderate-to-large systemic effect and do not evaluate local joint mechanisms.")

    set_paragraph(find_paragraph(doc, "Human transcriptomic studies pose four recurring problems."),
        "Human transcriptomic studies pose four recurring problems. First, case-control comparisons mix anatomical regions, disease severities and cell compositions. Second, single-cell tests can yield very small P values when cells rather than donors are treated as independent units. Third, programme weights obtained by constrained fitting are compositional: because they are rescaled to sum to one, an increase in one programme can arise purely from decreases in the others [23]. Fourth, single-gene expression and cell-state programmes may not travel together across tissues. The seven chondrocyte states described in human OA cartilage provide a reference for separating a continuous homeostatic-to-hypertrophic transition from the behaviour of any one transcript [9]. Independent paired cartilage data can then evaluate cross-dataset programme concordance [10,11].")
    set_paragraph(find_paragraph(doc, "Observational tissue data also cannot determine"),
        "Observational tissue data also cannot determine whether systemic modulation of PIEZO1 expression alters OA risk. Expression quantitative trait loci (eQTLs) can proxy long-term differences in gene expression for Mendelian randomization (MR), while colocalization tests whether exposure and outcome associations are compatible with a shared causal variant [12-15]. The present causal analysis uses a well-powered whole-blood eQTL resource; a null result bounds the systemic whole-blood instrument model but does not evaluate cartilage- or subchondral-bone-specific mechanisms.")
    set_paragraph(find_paragraph(doc, "We therefore evaluated four linked questions:"),
        "We therefore evaluated four linked questions: whether PIEZO1 is a disease-wide expression marker; whether within-donor joint compartment explains a stronger signal; whether PIEZO1-associated chondrocyte programmes show cross-dataset concordance once compositional constraints are respected; and whether whole-blood expression genetics support a systemic causal effect. The analysis was organised around donor-level and paired-tissue estimates; immune and regulatory extensions were retained as exploratory supplementary analyses.")

    p_study_design = find_paragraph(doc, "This secondary, discovery-driven reanalysis")
    set_paragraph(p_study_design,
        "This secondary, discovery-driven reanalysis used public human data. Dataset roles were defined for disease context, paired within-joint comparisons, donor-aware single-cell association, cross-dataset programme concordance and whole-blood expression genetics. Processing, quality control, final sensitivity definitions and eligibility screens for optional evidence are in Supplementary Methods S1-S12. The analysis plan was not preregistered; final-audit additions are identified.")

    set_paragraph(find_paragraph(doc, "We analysed seven GEO series"),
        "We analysed GSE51588 [32], GSE55235 [33], GSE82107 [34], GSE46750 [35], GSE104782 [9], GSE57218 [10] and GSE152805 [11], together with eQTLGen whole-blood cis-eQTL statistics and FinnGen R12 osteoarthritis outcomes. Dataset identifiers, deposited samples, exclusions, preprocessing and final analysis counts are in Supplementary Tables S1-S4 and Data S1.")

    p_axis_methods = find_paragraph(doc, "GSE104782 cells were assigned")
    set_paragraph(p_axis_methods,
        p_axis_methods.text.replace("CRTL1 was prespecified", "CRTL1 was planned") + " The paired Q1-Q4 Wilcoxon signed-rank sensitivity test used the asymptotic normal approximation with continuity correction.")

    set_paragraph(find_paragraph(doc, "State programmes were scored"),
        "State programmes were scored by one final standardized pipeline. GSE104782-derived top-100 state markers were fitted to gene-level, probe-collapsed bulk expression using non-negative L-BFGS-B weights followed by closure to sum one. Table S6 reports the raw closed-weight layer, while the log-ratio, zero-handling, additive-log-ratio and independent gene-set specifications in Table S13 use the same marker construction, expression preprocessing and fitted weights. The homeostatic-minus-pre-hypertrophic contrast, rather than any isolated coordinate, was the principal cross-dataset comparison.")

    set_paragraph(find_paragraph(doc, "In GSE152805, cells were assigned"),
        "Reference-internal validation was fully nested within GSE104782: for each held-out donor, top-100 state markers were reselected and centroids rebuilt from the other nine donors; 1,000 donor-preserving label permutations repeated the complete fold-specific pipeline. Full-reference centroids were then projected onto GEO-supplied sparse matrices from the six GSE152805 cartilage libraries by maximum unconstrained similarity. This external projection support is not external validation because independent gold-standard GSE152805 state labels were unavailable. States with held-out recall below 0.60 were excluded in a final-audit sensitivity analysis without reassigning cells; retained proportions were renormalized before symmetric Kitagawa decomposition. Assignment confidence used the top-1-minus-top-2 margin, and full-state within-state PIEZO1 required at least 50 cells per donor-compartment-state. These estimates address compartmental heterogeneity, not disease causality.")

    set_paragraph(find_paragraph(doc, "Five independent PIEZO1 cis-eQTLs"),
        "The three core MR assumptions were relevance (the cis-eQTLs predict whole-blood PIEZO1 expression), independence from exposure-outcome confounders, and exclusion restriction (the variants affect OA only through whole-blood PIEZO1 expression). F statistics assessed relevance; source-study adjustment, the non-overlapping two-sample design, heterogeneity, MR-Egger, leave-one-out, directionality and colocalization probed the other assumptions, which cannot be proven from summary data. Five independent PIEZO1 cis-eQTLs present in all outcomes were analysed. Because eQTLGen supplied Z scores and sample sizes, per-allele effects on standardized expression were reconstructed as beta=Z/sqrt[2f(1-f)N] and SE=1/sqrt[2f(1-f)N], using assessed-allele frequency f; F=Z^2. The primary estimate was inverse-variance weighted with multiplicative random effects; fixed-effect, DerSimonian-Laird, weighted-median and MR-Egger estimates were sensitivity checks. Variants absent from an outcome were excluded. No individual-level missing data were imputed, and no multiplicity adjustment was applied across three correlated OA endpoints.")

    set_paragraph(find_paragraph(doc, "Deposited age and sex were included"),
        "Deposited age and sex were included where available; missing covariates were not imputed. The RegC-FC contrast and other unplanned comparisons were labelled exploratory. Analyses were performed in R 4.6.1 and Python 3.12.13. Package versions, explicit model implementations and exact data-to-result links are listed in Supplementary Tables S16 and S17.")

    set_paragraph(find_paragraph(doc, "Colocalization used coloc.abf"),
        "Colocalization used coloc.abf in the PIEZO1 region and reported posterior probabilities for distinct (H3) and shared (H4) signals separately by outcome. Power boundaries used the primary multiplicative-random-effects standard errors and, separately, the more conservative DerSimonian-Laird standard errors at two-sided alpha=0.05 and 80% power.")
    set_paragraph(find_paragraph(doc, "The continuous axis was more reproducible"),
        "The continuous axis was more consistent across donors. Nine of ten donors showed a negative PIEZO1 correlation with the hypertrophic-minus-homeostatic axis, with a mean donor Spearman correlation of -0.137 (95% CI -0.216 to -0.058; one-sample t P=0.00346). Three donor-level sensitivity tests agreed: Fisher-z one-sample t P=0.00355 (mean z=-0.139), Wilcoxon signed-rank P=0.0108 (median rho=-0.156) and exact sign-test P=0.0215. With Q1 and Q4 defined within donor, PIEZO1 was higher at the homeostatic than hypertrophic end by 0.123 normalized-expression units (paired t P=0.0134; Wilcoxon P=0.0323). The cell-level association (rho=-0.142, P=5.26e-8) reflects 1,464 nested cells, explains 2.0% of variance and is descriptive. These results support a weak donor-consistent continuous state association, not a categorical ranking of all seven states (Figure 3; Supplementary Tables S5 and S20).")

    set_paragraph(find_paragraph(doc, "3.3 Donor-aware analysis"),
        "3.3 Donor-aware analysis supports a weak association of PIEZO1 with the homeostatic end of a chondrocyte continuum")

    set_paragraph(find_paragraph(doc, "We therefore asked whether the GSE152805"),
        "We therefore asked whether the GSE152805 compartment difference was an artefact of cell-state composition. Fully nested leave-one-donor-out validation, with marker selection repeated inside each training fold, gave balanced accuracy 0.595 and macro-F1 0.554, above the complete-pipeline donor-preserving permutation-null 95th percentile of 0.172 (P=0.001). Held-out recall was below 0.60 for HTC, RegC and EC. In the full-state analysis, symmetric within-state components were medial-higher in all three donors. Across donors and reference formulations, within-state shares occurred on both sides of 50%, and the donor 118 interval crossed 50%. In the recall-gated sensitivity retaining FC, HomC, ProC and preHTC without reassigning cells, total lateral-minus-medial differences remained negative in donors 113, 116 and 118 (-0.047, -0.138 and -0.069), but the within-state component was slightly positive in donor 113 (+0.004) and negative in donors 116 and 118 (-0.101 and -0.032); composition components were negative in all three (-0.051, -0.037 and -0.037). The medial PIEZO1 excess therefore persisted overall, but the relative contribution of within-state expression and composition was not uniform across donors (Figure 4; Supplementary Figures S12 and S15; Supplementary Tables S14, S19 and S21).")

    set_paragraph(find_paragraph(doc, "The apparent contradiction resolves"),
        "The apparent contradiction resolves once three quantities measured on different axes are separated. The first is the cross-state association: within GSE104782 donors, PIEZO1 declines weakly along a continuous hypertrophic-minus-homeostatic axis. The second is the projected within-state compartment contrast, which was predominantly medial-higher in the full-state analysis but was sensitive to the nested recall gate. The third is the compartment average, which combines within-state expression with cell-state composition. After retaining only states with recall at least 0.60, the total medial-higher direction persisted in all three donors, although donor 113 became composition dominated and its within-state component reversed slightly. These quantities are not mutually exclusive, and their different signs show that PIEZO1 transcript level is not a fixed function of position on the state axis.")

    set_paragraph(find_paragraph(doc, "3.4 The homeostatic-versus-pre-hypertrophic"),
        "3.4 The homeostatic-versus-pre-hypertrophic contrast shows cross-dataset concordance and survives compositional testing")
    set_paragraph(find_paragraph(doc, "Programme-level integration was secondary"),
        "Programme-level integration was secondary, so its evidential hierarchy is stated before the estimates. GSE104782 defines the programmes; GSE57218 provides paired-cartilage cross-dataset concordance; GSE152805 provides external projection support only; and GSE51588 contributes regional discovery and cross-tissue transcriptomic similarity. GSE152805 is not an external validation dataset because independent gold-standard state labels were unavailable.")
    set_paragraph(find_paragraph(doc, "The same contrast held in the subchondral-bone"),
        "The same contrast held in the subchondral-bone discovery dataset, where it is interpreted as cross-tissue transcriptomic resemblance rather than chondrocyte content: 0.6846 as log(HomC/preHTC) (P=8.2e-5) and 0.355 with the independent gene-set score (P=7.1e-5). Centred and additive log-ratio results paralleled GSE57218, with the pre-hypertrophic programme negative under every treatment (-0.222 to -1.163) and the homeostatic programme reference dependent (-0.479 to 0.685, P=0.052 to 0.37) and not significant with the independent gene-set score (0.057, P=0.35). GSE51588 lateral-up and medial-up signatures transferred to GSE57218 in the expected directions (0.115, P=0.0060 and -0.100, P=0.0140). In GSE152805, five of seven projected programmes had the same direction in all three donors; fibrocartilage and regulatory scores agreed in two. With three donors and no gold-standard query labels, this is external projection support, not independent statistical validation (Figure 4; Supplementary Figure S7).")
    set_paragraph(find_paragraph(doc, "3.5 PIEZO1 direction"),
        "3.5 PIEZO1 direction is context dependent, and the cartilage compartment effect is not purely compositional")
    set_paragraph(find_paragraph(doc, "PIEZO1 messenger RNA itself"),
        "PIEZO1 messenger RNA did not show one cross-dataset direction. Preserved cartilage in GSE57218 had lower PIEZO1 than OA-affected cartilage (preserved-minus-affected -0.253, P=0.013), and all three GSE152805 donors had lower mean PIEZO1 laterally than medially (-0.052, -0.147 and -0.073). Both directions oppose the subchondral-bone contrast, where PIEZO1 is higher laterally.")
    set_paragraph(find_paragraph(doc, "3.6 Blood-expression genetics"),
        "3.6 Whole-blood expression genetics bound the tested systemic instrument model")
    set_paragraph(find_paragraph(doc, "Five independent blood cis-eQTL instruments"),
        "Five independent whole-blood cis-eQTL instruments were available for all outcomes (F=32.7-654.6; all F>10). Under the primary multiplicative random-effects model, log odds ratios per standard-deviation higher genetically predicted whole-blood PIEZO1 expression were -0.027 (SE 0.026, P=0.30) for overall OA, -0.032 (SE 0.028, P=0.26) for knee OA and -0.005 (SE 0.032, P=0.87) for hip OA. Fixed-effect P values were 0.166, 0.179 and 0.874; DerSimonian-Laird, weighted-median and MR-Egger slope estimates were also nonsignificant (Figure 5; Supplementary Tables S10-S11).")
    set_paragraph(find_paragraph(doc, "Power is reported conservatively"),
        "Power is reported conservatively because the exposure scale is reconstructed. Using the most conservative variance model, minimum detectable odds ratios at 80% power were 1.14 for overall OA, 1.14 for knee OA and 1.09 for hip OA; under the primary model they were 1.08, 1.08 and 1.09. Power to detect an odds ratio of 1.10 was 55%, 54% and 85% under the conservative model and exceeded 97% for an odds ratio of 1.20 in all three outcomes. Within the whole-blood cis-eQTL instrument model, the estimates did not support a moderate-to-large systemic effect; they neither evaluate nor establish any local cartilage or subchondral-bone mechanism (Figure 5; Supplementary Table S12).")
    set_paragraph(find_paragraph(doc, "This analysis separates two quantities"),
        "This analysis separates two quantities that are easily conflated: a consistent cross-dataset shift in chondrocyte-programme balance within human OA joints, and the context-dependent expression of PIEZO1 itself. Paired-cartilage concordance came from GSE57218, external projection support from three GSE152805 donors, and cross-tissue regional similarity from GSE51588. The contrast survived compositional sensitivity testing. By contrast, PIEZO1 transcript abundance was higher laterally in subchondral bone but higher medially or in affected regions in two cartilage datasets. A single directional biomarker claim would discard this context dependence.")
    set_paragraph(find_paragraph(doc, "The genetic analysis supplies a boundary"),
        "The genetic analysis supplies a boundary rather than mechanistic proof, and its numerical scale depends on reconstructing per-allele eQTLGen effects. Uniform rescaling changes Wald-ratio and standard-error units proportionally but leaves the IVW z statistic, P value and heterogeneity statistics unchanged. Thus the null test is robust, whereas odds-ratio units and minimum detectable effects remain scale dependent. We report primary and conservative variance models with explicit scale-sensitivity bounds. These estimates apply only to the whole-blood instrument model and do not evaluate local joint-tissue mechanisms.")
    set_paragraph(find_paragraph(doc, "In conclusion, human OA tissues"),
        "In conclusion, human OA tissues show a consistent cross-dataset shift in the balance between homeostatic and pre-hypertrophic chondrocyte programmes that is robust to the compositional constraint on programme weights. PIEZO1 is embedded in this biology, but its transcript response is context dependent rather than universally directional. This framework motivates a focused next experiment: paired regional cartilage and subchondral-bone sampling with direct force, PIEZO1 activity and calcium measurements at donor level.")
    set_paragraph(find_paragraph(doc, "All transcriptomic datasets are cross-sectional"),
        "All transcriptomic datasets are cross-sectional and were generated on different platforms. The subchondral-bone discovery tissue and cartilage reference are biologically distinct, so programme weights in bone indicate cross-tissue transcriptomic resemblance, not cell fractions. GSE51588 included five non-OA donors and GSE152805 three paired OA donors; the latter offers external projection support rather than independent statistical validation. GSE104782 used deposited processed counts and annotations, whereas GSE152805 used GEO-supplied sparse matrices. No deposited doublet calls or ambient-RNA estimates were available, no additional correction was applied, and residual doublets or ambient RNA remain possible. Programme integration was secondary and discovery driven, not a preregistered primary endpoint. Body-mass index and radiographic grade were unavailable; residual within-donor confounding remains possible. Exploratory immune, transcription-factor, motif and drug-target analyses do not strengthen the central claims.")

    set_paragraph(find_paragraph(doc, "Decomposition also shows"),
        "Decomposition does not support a single universal mechanism for the cartilage compartment difference. In the full-state analysis, both within-state expression and composition were material. Across donors and reference formulations, within-state shares occurred on both sides of 50%, and the donor 118 interval crossed 50%. After excluding HTC, RegC and EC for nested recall below 0.60, both components aligned with the total difference in donors 116 and 118, whereas donor 113 was composition dominated with a small opposing within-state component. The result is donor-heterogeneous contributions, not a uniformly within-state or uniformly compositional mechanism.")

    set_paragraph(find_paragraph(doc, "Code availability."),
        "Code availability. A versioned local Git repository contains the analytical scripts, software versions, source results and data-to-result manifest used for this revision. Local release v1.1.0 is prepared for publication. A public repository URL and DOI-bearing archive require authenticated deposition and must be inserted before submission.")

    set_paragraph(find_paragraph(doc, "Ethics approval and consent"),
        "Ethics approval and consent to participate. This study was a secondary analysis of lawfully accessed, publicly available, de-identified datasets. It involved no new participant recruitment, intervention, specimen collection or access to identifiable private information. Article 32 of the Measures for Ethical Review of Life Science and Medical Research Involving Humans (National Health Commission, Ministry of Education, Ministry of Science and Technology, and National Administration of Traditional Chinese Medicine of China, 2023) permits exemption from ethics review for research using lawfully obtained public data or anonymized information when the specified conditions are met. Accordingly, no additional institutional ethics committee approval or participant consent was required for this secondary analysis. Ethics approval and informed consent for the original data collection were the responsibility of the original investigators and are reported in the source publications and database records.")
    set_paragraph(find_paragraph(doc, "Author contributions."),
        f"Author contributions. {AUTHOR}: Conceptualization; Data curation; Formal analysis; Investigation; Methodology; Project administration; Resources; Software; Supervision; Validation; Visualization; Writing, original draft; Writing, review and editing. {AUTHOR} approved the final manuscript and accepts full responsibility for the work.")
    set_paragraph(find_paragraph(doc, "Funding."),
        "Funding. This research received no specific grant from any funding agency in the public, commercial, or not-for-profit sectors. Funding for the source datasets and consortia is reported in their original publications and repository records; the author had no role in those awards.")
    set_paragraph(find_paragraph(doc, "Competing interests."),
        "Competing interests. The author declares no competing interests.")
    set_paragraph(find_paragraph(doc, "Acknowledgements."),
        "Acknowledgements. The author acknowledges the investigators and participants who generated and shared the GEO, eQTLGen, GTEx and FinnGen resources.")
    set_paragraph(find_paragraph(doc, "Use of generative artificial intelligence."),
        "Use of generative artificial intelligence. Generative AI tools were used to assist language editing, document structuring and consistency checking. All analyses, numerical values, citations and interpretations were reviewed against source outputs by the author, who retains full responsibility for the manuscript.")

    set_paragraph(find_paragraph(doc, "Figure 1."),
        "Figure 1. Evidence architecture and disease-level context. (a) Study design separating disease-level expression, paired joint compartments, donor-aware chondrocyte-state analyses, cross-dataset concordance, within-state decomposition and whole-blood expression genetics; the inferential unit of each layer is stated. (b) Dataset-specific PIEZO1 case-minus-control differences with 95% confidence intervals: subchondral bone -0.109 (-0.422 to 0.205; P=0.500), synovium GSE55235 +0.243 (-0.297 to 0.783; P=0.35) and synovium GSE82107 +0.841 (-0.054 to 1.736; P=0.063). Estimates remain on native scales and are not pooled. (c) Interpretation contract. The figure does not support PIEZO1 as a uniform disease-wide or standalone diagnostic marker.")
    set_paragraph(find_paragraph(doc, "Figure 4."),
        "Figure 4. Cross-dataset concordance of the chondrocyte-programme contrast and context-dependent PIEZO1 transcription. (a) Evidence hierarchy: reference definition, paired-cartilage concordance, external projection support and regional discovery. (b) Direction and robustness across raw, centred-log-ratio, additive-log-ratio, pairwise-log-ratio and independent gene-set specifications from the same final pipeline. The contrast remained positive under every final transformation, zero-handling rule and reference specification in both datasets; the isolated homeostatic increase was reference dependent. (c) GSE152805 projected programme directions by donor. (d) Full-reference within-state lateral-minus-medial PIEZO1 differences. (e) Full-reference symmetric decomposition into donor-heterogeneous within-state expression and composition components; diamonds mark totals. (f) Dataset-specific PIEZO1 directions. Effect magnitudes are not compared across transformations or datasets; the final-audit recall-gated sensitivity is shown in Supplementary Figure S15.")
    set_paragraph(find_paragraph(doc, "Figure 5."),
        "Figure 5. Whole-blood expression genetic evidence. (a) IVW estimates under three variance models; the outline marks the primary multiplicative model. (b) Weighted-median and MR-Egger sensitivity estimates. (c) Per-SNP Wald ratios and IVW weights. (d) Colocalization posterior probabilities for distinct and shared signals. (e) Minimum detectable odds ratios at 80% power. Bars in panels a-c are 95% confidence intervals. Inference is limited to the whole-blood cis-eQTL instrument model and does not evaluate local joint mechanisms.")

    references = [
        "32. Chou CH, Wu CC, Song IW, et al. Genome-wide expression profiles of subchondral bone in osteoarthritis. Arthritis Res Ther. 2013;15:R190. doi:10.1186/ar4380.",
        "33. Woetzel D, Huber R, Kupfer P, et al. Identification of rheumatoid arthritis and osteoarthritis patients by transcriptome-based rule set generation. Arthritis Res Ther. 2014;16:R84. doi:10.1186/ar4526.",
        "34. Broeren MGA, de Vries M, Bennink MB, et al. Functional tissue analysis reveals successful cryopreservation of human osteoarthritic synovium. PLoS One. 2016;11:e0167076. doi:10.1371/journal.pone.0167076.",
        "35. Lambert C, Dubuc JE, Montell E, et al. Gene expression pattern of cells from inflamed and normal areas of osteoarthritis synovial membrane. Arthritis Rheumatol. 2014;66:960-968. doi:10.1002/art.38315.",
        "36. Katsoula G, Arruda AL, Tutino M, et al. Transcriptome analysis in osteoarthritis primary tissues identifies high-confidence effector genes. Nat Commun. 2026;17:8128. doi:10.1038/s41467-026-74993-y.",
    ]
    anchor = find_paragraph(doc, "Figure legends")
    reference_style = find_paragraph(doc, "31. Higgins").style
    for reference in references:
        anchor.insert_paragraph_before(reference, style=reference_style)

    replace_image_by_order(doc, 0, FIG1)
    replace_image_by_order(doc, 3, FIG4)
    replace_image_by_order(doc, 4, FIG5)

    abstract_text = " ".join(find_paragraph(doc, label).text for label in ("Objective.", "Design.", "Results.", "Conclusions."))
    intro_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text == "1 Introduction")
    decl_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text == "Declarations")
    main_text = " ".join(p.text for p in doc.paragraphs[intro_idx + 1:decl_idx])
    abstract_words = count_words(abstract_text)
    main_words = count_words(main_text)
    set_paragraph(find_paragraph(doc, "Manuscript type:"),
        f"Manuscript type: Full-length original research article | Main-text word count: {main_words:,} | Abstract word count: {abstract_words}")
    if abstract_words > 340:
        raise RuntimeError(f"Abstract exceeds gate: {abstract_words}")
    if main_words >= 4000:
        raise RuntimeError(f"Main text exceeds gate: {main_words}")

    clean_properties(doc, doc.paragraphs[0].text,
                     "Submission-audit HOLD copy following the corrected V11/V9 final-review SOP")
    return doc


def revise_supplement() -> Document:
    doc = Document(SOURCE_SUPP)

    set_paragraph(find_paragraph(doc, "[Author names"),
        f"{AUTHOR} (ORCID: {ORCID})")

    set_paragraph(find_paragraph(doc, "This supplement separates"),
        "This supplement separates confirmatory, supportive and exploratory analyses. Every numerical statement in the main manuscript is traceable to a stored result file, and the script that produced each result is listed in Supplementary Table S17. Tables S13 and S14 report compositional sensitivity and within-state decomposition; Table S18 records the final eligibility decisions for optional comparator, spatial and joint-tissue genetic extensions.")
    set_paragraph(find_paragraph(doc, "GSE51588 was the paired regional discovery dataset"),
        "GSE51588 was the paired regional discovery dataset. GSE104782 defined chondrocyte states and enabled donor-aware analysis. GSE57218 supplied paired-cartilage cross-dataset programme concordance. GSE152805 supplied compartment-matched external projection support and decomposition, not independent validation. Synovial datasets tested disease-wide direction and exploratory immune context. Genetic resources were not treated as tissue-concordance datasets. Subchondral-bone programme weights indicate cross-tissue transcriptomic similarity, not chondrocyte content.")
    set_paragraph(find_paragraph(doc, "The reproducible signal in osteoarthritis"),
        "The most consistent signal in osteoarthritis is a shift in the balance between homeostatic and pre-hypertrophic chondrocyte programmes, reported as the homeostatic-minus-pre-hypertrophic contrast rather than as uniform cross-tissue direction of PIEZO1 transcription. Whole-blood expression genetics constrain the size of a systemic effect but cannot evaluate local joint mechanisms. The table below fixes the terms used throughout the manuscript and supplement.")

    p_projection = find_paragraph(doc, "For each GSE104782 state")
    set_paragraph(p_projection,
        "For each GSE104782 state, fold changes were computed against the mean of the remaining states and the top 100 genes were retained after requiring mean normalized expression above 0.05, giving 690 unique signature genes. The authoritative final bulk programme pipeline converted the GSE104782 normalized matrix back to the linear scale for centroid construction; bulk probes were collapsed to gene symbols by their mean, converted from log2 to linear scale when appropriate, and standardized gene-wise across samples. Non-negative L-BFGS-B weights were fitted and then closed to sum one. Table S6 reports this L0 raw-weight layer, and Table S13 applies all zero-handling, log-ratio and independent gene-set sensitivity layers to the same final marker construction and bulk preprocessing. For GSE152805, query cells and reference centroids were standardized over frozen genes and assigned by the highest unconstrained similarity projection score; these scores are arbitrary units, not proportions. PIEZO1 was absent from all top-100 marker sets.")

    p_axis = find_paragraph(doc, "Continuous axis evidence is primary")
    set_paragraph(p_axis, p_axis.text.replace("CRTL1 was prespecified", "CRTL1 was planned") +
        " The paired Q1-Q4 Wilcoxon signed-rank sensitivity used the asymptotic normal approximation with continuity correction.")

    p_mr_source = find_paragraph(doc, "The deposited eQTLGen PIEZO1 cis file")
    set_paragraph(p_mr_source, p_mr_source.text +
        " The exposure and FinnGen R12 outcome samples were treated as non-overlapping. FinnGen registry-derived endpoints were overall OA (M13_ARTHROSIS), knee OA (gonarthrosis) and hip OA (coxarthrosis); case and control counts are reported in Supplementary Table S10. Individual-level participant data and covariates were unavailable to this secondary analysis.")

    set_paragraph(find_paragraph(doc, "Three inverse-variance-weighting variance models"),
        "Three inverse-variance-weighting variance models were computed. The fixed-effect model used weights 1/se^2; the multiplicative random-effects model inflated its standard error by sqrt(max(1, Q/df)); and the DerSimonian-Laird model re-estimated tau^2 and refitted with weights 1/(se^2+tau^2). The multiplicative model was designated as the primary final model; the others were retained because variance estimators diverge when one instrument dominates the weights. Cochran Q, heterogeneity P and I-squared came from the fixed-effect fit. MR-Egger, weighted median, leave-one-out and approximate Steiger analyses followed the implementations listed in Table S16.")
    set_paragraph(find_paragraph(doc, "Immune deconvolution, transcription-factor activity"),
        "Immune deconvolution, transcription-factor activity, motif enrichment, PROGENy, transcription-factor MR, transcriptome-wide association, fine-mapping and drug-target analyses are retained for transparency but cannot strengthen the central causal claim because donor-level independence, tissue matching, linkage pruning or independent corroboration was incomplete for at least one component.")
    set_paragraph(find_paragraph(doc, "S12 Prespecified eligibility screening"),
        "S12 Final eligibility screening for optional external evidence")
    set_paragraph(find_paragraph(doc, "Before inspecting optional extension results"),
        "A fixed comparator panel comprised PIEZO1, PIEZO2, TRPV4, TRPV2 and TRPA1. Eligible comparison required donor-level estimands across GSE51588, GSE104782, GSE57218 and GSE152805 without result-dependent substitution. Spatial eligibility required human OA cartilage or osteochondral tissue, at least three independent donors, identifiable donor labels, a defined spatial axis and donor-level inference. Joint-tissue QTL eligibility required human joint tissue and PIEZO1 cis-region statistics with effect allele, beta, standard error, ancestry or linkage-disequilibrium information, sample size and method. Failure of an optional screen was recorded as not eligible for primary analysis and did not alter manuscript claims.")
    set_paragraph(find_paragraph(doc, "Cell-level subtype and axis tests reproduced"),
        "Cell-level subtype and axis tests gave small P values, but donor-level analyses changed the conclusion. The categorical state omnibus was not significant by Friedman test, whereas the continuous axis was negative in nine of ten donors under all four donor-level sensitivity tests. With Q1/Q4 defined within donor, the mean Q1-minus-Q4 PIEZO1 difference was 0.123 normalized-expression units (paired t P=0.0134; Wilcoxon P=0.0323). Pseudobulk provided a nominal omnibus signal that did not survive transcriptome-wide correction; one pairwise contrast, RegC versus FC, survived and remains exploratory because it followed a non-significant omnibus test.")
    set_paragraph(find_paragraph(doc, "The prespecified comparator assessment"),
        "The final comparator assessment covered only two of the four required datasets and omitted two fixed genes, so it was not eligible for primary analysis and could not support a complete channel-specificity claim. GSE254844 provides human OA cartilage Geo-seq profiles across histological zones and weight-bearing regions, but public metadata do not support independent donor-level zone inference. A 2026 OA joint-tissue eQTL resource included several joint tissues, but PIEZO1 was absent from its packaged significant and conditionally independent cis-eQTL tables. No eligible local PIEZO1 instrument was available. These extensions were not eligible for primary analysis; the main narrative remains bounded to donor-level transcriptomics and whole-blood genetics.")

    set_paragraph(find_paragraph(doc, "Supplementary Table S6."),
        "Supplementary Table S6. Cross-dataset programme and transferred-signature concordance from the authoritative final pipeline")
    set_paragraph(find_paragraph(doc, "Supplementary Table S14."),
        "Supplementary Table S14. Within-state PIEZO1 and donor-heterogeneous Kitagawa contributions in GSE152805 (lateral minus medial)")
    set_paragraph(find_paragraph(doc, "Supplementary Table S18."),
        "Supplementary Table S18. Final optional-evidence eligibility audit")
    set_paragraph(find_paragraph(doc, "Supplementary Table S19."),
        "Supplementary Table S19. Reference-internal state-transfer validation and external-projection confidence sensitivity")
    set_paragraph(find_paragraph(doc, "Supplementary Table S21."),
        "Supplementary Table S21. Symmetric and reference-sensitive Kitagawa decomposition showing donor-heterogeneous contributions")
    set_paragraph(find_paragraph(doc, "Supplementary Figure S6."),
        "Supplementary Figure S6. Programme and transferred-signature concordance. (a) Full paired programme-similarity estimates for GSE51588 and GSE57218 with 95% confidence intervals. (b) Transfer of lateral-up and medial-up compartment signatures to GSE57218 preserved-minus-affected cartilage scores (P=0.0060 and P=0.0140). Similarity weights are not cell fractions. The authoritative final programme estimates and compositional sensitivity are reported in Tables S6 and S13 and Supplementary Figure S11.")
    set_paragraph(find_paragraph(doc, "Supplementary Figure S8."),
        "Supplementary Figure S8. Mendelian-randomization sensitivity analyses. Per-SNP Wald ratios with 95% confidence intervals for overall OA, leave-one-out inverse-variance-weighted estimates, IVW weight and approximate Steiger directionality. These analyses use reconstructed whole-blood PIEZO1 cis-eQTL instruments and are bounded to that systemic instrument model; they do not evaluate local joint-tissue mechanisms.")

    set_paragraph(find_paragraph(doc, "Raw 10x profiles were retained"),
        "GSE104782 used the deposited processed count matrix and accompanying cell-quality and clustering annotations. GSE152805 used GEO-supplied matrix.mtx, barcodes.tsv and genes.tsv files for six cartilage libraries; three synovial libraries were excluded before analysis. Cells required at least 200 detected genes, at least 500 counts and mitochondrial fraction below 25%; all 26,228 deposited cartilage barcodes passed, so no library-specific exclusion was introduced. Counts were log-normalized to 10,000 per cell. Deposited doublet calls and ambient-RNA estimates were unavailable; no additional doublet removal or ambient-RNA correction was applied, and the supplied barcode matrices lacked empty droplets for a de novo ambient profile. Residual doublets or ambient RNA therefore remain limitations. Reference-internal validation was fully nested within GSE104782, with fold-specific marker selection and 1,000 complete-pipeline donor-preserving permutations. Full-reference centroids were then projected into GSE152805 as external projection support, not external validation, because no independent gold-standard query-state labels were available. States with held-out recall below 0.60 were excluded only in the final-audit sensitivity analysis without reassigning cells. Full-state within-state PIEZO1 required at least 50 cells per donor-compartment-state; symmetric Kitagawa decomposition and margin-trimming sensitivities followed the stated rules.")

    set_paragraph(find_paragraph(doc, "Within-state results are given"),
        "Within-state results are given in Supplementary Table S14. Reference-internal fully nested leave-one-donor-out validation gave balanced accuracy 0.595 and macro-F1 0.554, above the complete-pipeline donor-preserving permutation-null 95th percentile of 0.172 (P=0.001). HTC, RegC and EC had recall below 0.60 and were excluded only in the final-audit sensitivity. In the full-state analysis, symmetric within-state shares were 28.7%, 77.2% and 50.5%, spanning both sides of 50% across donors and reference formulations. After retaining FC, HomC, ProC and preHTC without reassigning cells, total lateral-minus-medial differences remained negative in all donors (-0.047, -0.138 and -0.069), but donor 113 had a small positive within-state component (+0.004) opposing a negative composition component (-0.051); both components were negative in donors 116 and 118. The result is donor-heterogeneous contributions and does not support a uniform mechanism (Supplementary Tables S19 and S21; Supplementary Figure S15).")

    set_paragraph(find_paragraph(doc, "Supplementary Figure S12."),
        "Supplementary Figure S12. Within-state expression and compositional contributions in GSE152805 using full-reference projected states. (a) Lateral-minus-medial PIEZO1 within projected state and donor; symbol area reflects the smaller compartment cell count. (b) Lateral-minus-medial state proportions. (c) Symmetric Kitagawa decomposition into within-state expression and composition components; diamonds mark total differences. Contributions were donor heterogeneous and do not support a uniform mechanism. The final-audit recall-gated sensitivity is shown in Supplementary Figure S15 and Supplementary Table S21.")

    set_paragraph(find_paragraph(doc, "Supplementary Figure S15."),
        "Supplementary Figure S15. Reference-internal fully nested validation and external projection sensitivity analyses. (a) Leave-one-donor-out confusion matrix within GSE104782, with marker selection repeated in each training fold. (b) Observed balanced accuracy against 1,000 donor-preserving complete-pipeline permutations. (c) Held-out recall; the dashed line marks the 0.60 threshold used only in the final-audit sensitivity analysis. (d) GSE152805 projection-score margins, reported as external projection support rather than external validation. (e) Symmetric decomposition after excluding HTC, RegC and EC without reassigning cells. The total medial-higher direction persisted, but donor 113 was composition dominated with a small opposing within-state component.")

    # Supplementary Table S1: clarify the complete deposited sample sets.
    glossary = doc.tables[0]
    set_cell(glossary.rows[7].cells[2], "With three donors this is external projection support, not independent statistical validation")
    set_cell(glossary.rows[8].cells[2], "Bounded to the whole-blood instrument model; does not evaluate local cartilage or bone mechanisms")

    t1 = doc.tables[1]
    set_cell(t1.rows[2].cells[3], "30 deposited: 10 OA + 10 healthy used; 10 RA excluded")
    set_cell(t1.rows[3].cells[3], "17 deposited: 10 OA + 7 healthy; all used")
    set_cell(t1.rows[6].cells[4], "Paired-cartilage cross-dataset programme concordance")
    set_cell(t1.rows[7].cells[4], "External projection support; within-state decomposition")

    # Supplementary Table S1A: exact deposited, excluded and final sample flow.
    t2 = doc.tables[2]
    for row in t2.rows[1:]:
        if row.cells[0].text == "GSE55235":
            values = ["GSE55235", "30 arrays", "10 RA arrays", "RA comparator outside OA-versus-healthy analysis", "20", "20", "none"]
            for cell, value in zip(row.cells, values): set_cell(cell, value)
        if row.cells[0].text == "GSE82107":
            values = ["GSE82107", "17 arrays", "0", "-", "17", "17", "none"]
            for cell, value in zip(row.cells, values): set_cell(cell, value)

    # Supplementary Table S2: reflect donor heterogeneity after the recall gate.
    t3 = doc.tables[3]
    for row in t3.rows[1:]:
        if row.cells[0].text == "The homeostatic-minus-pre-hypertrophic programme contrast replicates":
            values = [
                "The homeostatic-minus-pre-hypertrophic programme contrast is concordant across datasets",
                "Supported within the analysed datasets",
                "Same final pipeline; five analytic layers in two datasets; three-donor external projection support",
                "Cross-dataset concordance, not independent external validation",
            ]
            for cell, value in zip(row.cells, values): set_cell(cell, value)
        if row.cells[0].text == "PIEZO1 transcript direction replicates":
            values = [
                "PIEZO1 transcript direction is uniform across datasets",
                "Not supported",
                "Bone and cartilage directions differ",
                "Context dependence is central",
            ]
            for cell, value in zip(row.cells, values): set_cell(cell, value)
        if row.cells[0].text == "The cartilage compartment effect is purely compositional":
            values = [
                "The cartilage compartment effect is purely compositional",
                "Not supported overall; donor heterogeneity",
                "Recall-gated total direction persisted; within-state and composition components were concordant in 2/3 donors",
                "Donor 113 was composition dominated; do not claim a uniform mechanism",
            ]
            for cell, value in zip(row.cells, values): set_cell(cell, value)

    # Supplementary Table S6: use the L0 layer from the same final programme pipeline as Table S13/Figure 4.
    t7 = doc.tables[7]
    s6_rows = {
        1: ["GSE51588", "Lateral - medial", "HomC", "+0.0356", "P 4.94e-4; FDR 0.00173", "Homeostatic higher laterally"],
        2: ["GSE51588", "Lateral - medial", "RegC", "+0.0446", "P 0.00767; FDR 0.0134", "Regulatory higher laterally"],
        3: ["GSE51588", "Lateral - medial", "preHTC", "-0.0612", "P 3.77e-4; FDR 0.00173", "Pre-hypertrophic higher medially"],
        4: ["GSE57218", "Preserved - affected", "HomC", "+0.0298", "P 4.37e-5; FDR 6.11e-5", "Homeostatic higher preserved"],
        5: ["GSE57218", "Preserved - affected", "preHTC", "-0.1197", "P 1.78e-7; FDR 6.22e-7", "Pre-hypertrophic higher affected"],
        6: ["GSE57218", "Preserved - affected", "GSE51588 lateral-up signature", "+0.1147", "P 0.0060", "Transferred direction concordant"],
        7: ["GSE57218", "Preserved - affected", "GSE51588 medial-up signature", "-0.0997", "P 0.0140", "Transferred direction concordant"],
    }
    for row_idx, values in s6_rows.items():
        for cell, value in zip(t7.rows[row_idx].cells, values): set_cell(cell, value)
    # Supplementary Table S5: authoritative Wilcoxon row.
    t6 = doc.tables[6]
    qrow = t6.rows[6]
    set_cell(qrow.cells[2], "V=49; median difference 0.1412; asymptotic with continuity correction")
    set_cell(qrow.cells[3], "0.0323")
    set_cell(t6.rows[11].cells[4], "Descriptive; cells are not independent inferential units")

    # Supplementary Table S12: primary versus conservative MR model terminology.
    t13 = doc.tables[13]
    for row in t13.rows[1:]:
        if row.cells[0].text == "MDE odds ratio at 80% power, prespecified model":
            set_cell(row.cells[0], "MDE odds ratio at 80% power, primary model")

    # Supplementary Table S16: complete software and explicit base implementations.
    t18 = doc.tables[18]
    additions = [
        ["Matrix", "1.7.5", "Sparse matrices"],
        ["coloc", "5.2.3", "Colocalization"],
        ["svglite", "2.2.2", "SVG export"],
        ["ragg", "1.5.2", "PNG/TIFF export"],
        ["scales", "1.4.0", "Plot-axis formatting"],
        ["stats::optim", "R 4.6.1 base", "Non-negative L-BFGS-B fitting, maxit=3000, followed by closure to sum one"],
        ["Base-R log transforms", "R 4.6.1 base", "Explicit CLR/ALR/log-ratio formulae; no external compositional package"],
        ["Python", "3.12.13", "MR power, GTEx scale validation and document QA"],
        ["numpy", "2.3.5", "Python numerical checks"],
        ["pandas", "3.0.1", "Python tabular checks"],
        ["python-docx", "1.2.0", "DOCX assembly and QA"],
        ["Pillow", "12.3.0", "Image inspection"],
    ]
    for values in additions:
        append_styled_row(t18, values)

    # Supplementary Table S17: no wildcards; every authoritative result is explicit.
    t19 = doc.tables[19]
    exact_results = {
        "GSE51588 paired, disease, interaction": "results/07_DEG_paired_MedialVsLateral_OA.csv; results/07_DEG_interaction.csv; results/07_DEG_disease_OAvsNormal.csv",
        "GSE104782 donor-aware inference": "results/40_P0_axis_donorlevel_tests.csv; results/40_P0_cell_vs_donor_comparison.csv; results/40_P0_cells_per_donor_subtype.csv; results/40_P0_donor_x_subtype_piezo1.csv; results/40_P0_donorlevel_axis_perdonor.csv; results/40_P0_donorlevel_subtype_omnibus.csv; results/40_P0_donorlevel_subtype_paired.csv; results/40_P0_pseudobulk_piezo1_contrasts.csv; results/40_P0_pseudobulk_piezo1_omnibus.csv",
        "Programme scoring, GSE51588 and GSE57218": "results/35_replication_signatures.csv",
        "Compositional sensitivity": "results/51_compositional_sensitivity_full.csv; results/51_compositional_sensitivity_key.csv; results/51_weights_GSE51588.csv; results/51_weights_GSE57218.csv",
        "GSE152805 donor directions": "results/39_gse152805_perdonor.csv; results/39_gse152805_piezo1.csv; results/39_gse152805_programmes.csv",
        "GSE152805 within-state and decomposition": "results/52_kitagawa_decomposition.csv; results/52_overall_direction.csv; results/52_state_composition.csv; results/52_within_state_per_donor.csv",
        "MR exposure audit, reconstruction and main results": "results/50_instrument_scale_audit.csv; results/50_mr_rerun.csv; results/50_mr_heterogeneity.csv; results/50_mr_ivw_weights.csv; results/50_mr_leave_one_out.csv; results/50_mr_steiger.csv; results/50_mr_scale_invariance.csv; results/50_mr_power.csv; results/50_mr_power_scale_sensitivity.csv; results/50_gtex_scale_validation.csv; results/50_gtex_scale_validation_summary.txt",
        "PIEZO1 colocalization": "results/11_coloc_PP.csv; results/11_coloc_region_data.csv; results/41_P0_coloc_PPH3_report.csv",
        "MR power": "results/32_power_primary_mr.csv",
        "Covariate audit and sensitivity": "results/53_covariate_audit_models.csv; results/53_covariate_audit.csv; results/54_covariate_sensitivity.csv",
        "State-transfer validation and confidence trimming": "04_Nested_LODO_QC.csv; 04_LODO_confusion_matrix.csv; 04_Nested_LODO_CellPredictions.csv; 04_Nested_LODO_Markers.csv; 04_Nested_LODO_PermutationNull.csv; 05_HighRecall_State_Sensitivity.csv; 05_HighRecall_State_Directions.csv; Figures/Figure_S15_Nested_StateTransfer_Validation.svg",
    }
    script_updates = {
        "Programme scoring, GSE51588 and GSE57218": "scripts/35_replication_gse57218.R (transferred signatures only)",
        "Compositional sensitivity": "scripts/51_P1_compositional_sensitivity.R (authoritative final programme pipeline)",
        "State-transfer validation and confidence trimming": "scripts/80_final_execution_nested_validation.R; scripts/83_final_audit_tables.R",
        "V10 document transformation": "scripts/84_apply_final_execution_docx.py",
    }
    for row in t19.rows[1:]:
        claim = row.cells[0].text
        if claim in exact_results:
            set_cell(row.cells[2], exact_results[claim])
        if claim in script_updates:
            set_cell(row.cells[1], script_updates[claim])
        if claim == "V10 document transformation":
            set_cell(row.cells[0], "Final execution document transformation")
            set_cell(row.cells[2], "PIEZO1_OA_Main_Manuscript_V13_Corrected_SOP_Audit_HOLD.docx; PIEZO1_OA_Supplementary_Materials_V11_Corrected_SOP_Audit_HOLD.docx; PIEZO1_OA_STROBE-MR_Checklist_Completed_V2.docx")
        if claim == "Programme scoring, GSE51588 and GSE57218":
            set_cell(row.cells[0], "Transferred compartment signatures")
        if claim == "Compositional sensitivity":
            set_cell(row.cells[0], "Programme scoring and compositional sensitivity (authoritative final pipeline)")
    append_styled_row(t19, ["Figure 1 P-value correction", "scripts/82_rebuild_figure1_final.R", "Figures/Figure1_Final_P0500.svg; Figures/Figure1_Final_P0500.pdf; Figures/Figure1_Final_P0500.png; Figures/Figure1_Final_P0500.tiff"])
    append_styled_row(t19, ["GEO sample-flow audit and Data S1", "scripts/83_final_audit_tables.R", "03_GEO_SampleFlow_Audit.csv; Data_S1_GSM_Donor_Region_Mapping.csv"])

    # Supplementary Table S18: cite the formal 2026 publication.
    t20 = doc.tables[20]
    for row in t20.rows[1:]:
        for cell in row.cells:
            if "Prespecified" in cell.text or "prespecified" in cell.text:
                set_cell(cell, cell.text.replace("Prespecified", "Fixed").replace("prespecified", "defined"))
        if row.cells[0].text == "Joint-tissue genetics":
            set_cell(row.cells[2], "OA joint-tissue eQTL map [36]")
            set_cell(row.cells[5], "Whole-blood MR is bounded to its systemic instrument model; it neither evaluates nor establishes local mechanisms")
        if row.cells[0].text == "Spatial validation":
            set_cell(row.cells[5], "No independent spatial-validation claim")

    # Supplementary Table S19: nested metrics and explicit state-recall gate.
    t21 = doc.tables[21]
    metrics = {
        "balanced_accuracy": ("0.595", "0.172", "PASS; complete-pipeline permutation P=0.001"),
        "macro_F1": ("0.554", "-", "descriptive"),
        "recall_ProC": ("0.703", "-", "PASS (>=0.60)"),
        "recall_HomC": ("0.821", "-", "PASS (>=0.60)"),
        "recall_HTC": ("0.454", "-", "FAIL (<0.60); excluded in sensitivity"),
        "recall_RegC": ("0.230", "-", "FAIL (<0.60); excluded in sensitivity"),
        "recall_preHTC": ("0.747", "-", "PASS (>=0.60)"),
        "recall_FC": ("0.673", "-", "PASS (>=0.60)"),
        "recall_EC": ("0.535", "-", "FAIL (<0.60); excluded in sensitivity"),
    }
    for row in t21.rows[1:]:
        key = row.cells[0].text
        if key in metrics:
            for cell, value in zip(row.cells[1:], metrics[key]): set_cell(cell, value)
    set_cell(t21.rows[10].cells[3], "Full-state within-state component remained medial-higher after margin trimming")
    append_styled_row(t21, ["Recall-gated sensitivity", "Retained FC, HomC, ProC and preHTC", "-", "Total medial-higher direction persisted in 3/3 donors; both components concordant in 2/3"])

    # Supplementary Table S21: exact reference terminology plus gated rows.
    t23 = doc.tables[23]
    set_cell(t23.rows[0].cells[5], "Lateral-/medial-reference within-state share range")
    high_recall_rows = [
        ["113", "High-recall states only", "-0.0469", "+0.0036", "-0.0505", "-20.8 to 5.4%", "2.8e-17"],
        ["116", "High-recall states only", "-0.1382", "-0.1013", "-0.0369", "72.6 to 74.0%", "2.8e-17"],
        ["118", "High-recall states only", "-0.0694", "-0.0321", "-0.0373", "22.4 to 70.2%", "2.8e-17"],
    ]
    for values in high_recall_rows:
        append_styled_row(t23, values)

    # Planned-marker wording without implying preregistration.
    t22 = doc.tables[22]
    for row in t22.rows[1:]:
        for cell in row.cells:
            if "full prespecified set" in cell.text:
                set_cell(cell, cell.text.replace("full prespecified set", "full planned marker set"))

    # Remove the internal author-completion page from the submission-facing supplement.
    remove_paragraph(find_paragraph(doc, "1. Confirm author names"))
    remove_paragraph(find_paragraph(doc, "Items requiring author completion before submission"))

    replace_image_by_order(doc, 0, FIGS1)
    replace_image_by_order(doc, -1, FIGS15)
    clean_properties(doc, "PIEZO1 OA Supplementary Materials",
                     "Submission-audit HOLD supplement following the corrected V11/V9 final-review SOP")
    return doc


def build_strobe_mr_checklist() -> Document:
    doc = Document(STROBE_TEMPLATE)
    set_paragraph(doc.paragraphs[1],
        "Manuscript: Human osteoarthritis reveals context-dependent PIEZO1 expression across chondrocyte-state programmes | Author: Yang Yunze | Completed: 1 September 2026")
    table = doc.tables[0]
    entries = {
        1: ("Main 1", "Abstract Design identifies Mendelian randomization; Results reports the MR estimates and power boundary."),
        3: ("Main 2", "Introduction explains biological plausibility, why observational tissue data cannot establish systemic causality, and why MR is informative."),
        4: ("Main 1-2", "Abstract Objective and final Introduction paragraph state the blood-expression genetic question; MR interpretation is explicitly conditional on IV assumptions."),
        6: ("Main 3-4; Supp 3-4", "Study design, public data sources, exposure reconstruction, outcomes, colocalization and power methods."),
        7: ("Main 3-4; Supp 3-4", "eQTLGen whole-blood exposure and FinnGen R12 OA outcomes; source setting and summary-data design are stated."),
        8: ("Main 3-4, 7; Supp 9-11", "Instrument availability, endpoint case/control counts, exclusions and power calculations are reported."),
        9: ("Main 4; Supp 3, 9-10", "cis-window selection, genome-wide significance, LD pruning, harmonization, effect reconstruction and F statistics."),
        10: ("Main 3-4; Supp 3, 10", "Whole-blood PIEZO1 expression is the exposure; FinnGen overall, knee and hip OA registry endpoints are the outcomes."),
        11: ("Main 9", "Ethics statement explains the public, de-identified secondary analysis and the Article 32 exemption basis; original-study approvals remain with source investigators."),
        12: ("Main 4; Supp 3-4", "Relevance, independence and exclusion restriction are stated; assumptions for sensitivity estimators are described."),
        13: ("Main 4; Supp 3-4", "IVW multiplicative random effects is primary; fixed-effect, DerSimonian-Laird, weighted median, MR-Egger and diagnostic methods are specified."),
        14: ("Main 4; Supp 3, 10-11", "Effects are per 1-SD higher genetically predicted whole-blood PIEZO1 expression; OR, SE, CI and power scales are defined."),
        15: ("Main 4; Supp 3, 9-10", "Independent cis-eQTLs were harmonized; IVW weights and dominant-instrument concentration are reported."),
        16: ("Main 4; Supp 3-4", "Wald-ratio/IVW construction, variance models and source-study covariate handling are described."),
        17: ("Main 4", "No individual-level data were available for imputation; variants absent from an outcome dataset were excluded."),
        18: ("Main 4", "No multiplicity adjustment across three correlated OA endpoints; exact P values and uncertainty are reported."),
        19: ("Main 4, 7-9; Supp 3-4, 10", "F statistics, MR-Egger, heterogeneity, leave-one-out, Steiger directionality and colocalization assess assumptions."),
        20: ("Main 4, 7-9; Supp 3-4, 10-11", "Alternative estimators/variance models, leave-one-out, scale sensitivity, colocalization and power analyses."),
        22: ("Main 4; Supp 12-14", "R 4.6.1, Python 3.12.13, package versions, settings and exact data-to-result map."),
        23: ("Main 3", "The study protocol and analysis plan were not preregistered."),
        26: ("Main 3, 7; Supp 6-7, 9-10", "Dataset/sample flow, five retained instruments and the excluded outcome-missing variant are reported."),
        27: ("Main 7; Supp 9-10", "Instrument strength and case/control counts are reported; individual-level phenotype summaries were unavailable for this summary-data MR."),
        28: ("Main 7; Supp 10", "Cochran Q, heterogeneity P and I-squared are reported across genetic instruments; source meta-analysis heterogeneity was unavailable."),
        29: ("Main 4; Supp 3", "Exposure and outcome samples were treated as non-overlapping; source-study adjustment and ancestry scope are stated."),
        31: ("Supp 9-10", "Variant-exposure beta/SE/F and variant-outcome-derived Wald/weight diagnostics are reported."),
        32: ("Main 7; Supp 10", "MR estimates with SE and P values are reported per 1-SD higher genetically predicted expression; Figure 5 provides 95% CIs."),
        33: ("N/A", "Absolute-risk translation was not performed because the estimand is a lifelong genetic proxy, the estimates were null and follow-up time was not defined."),
        34: ("Main 17", "Figure 5 presents forest-style IVW, sensitivity-estimator and per-SNP results."),
        36: ("Main 7-9; Supp 10", "Instrument strength, pleiotropy, directionality, dominant weight and lack of colocalization are reported."),
        37: ("Main 7; Supp 10", "Cochran Q, heterogeneity P, I-squared, MR-Egger intercept and IVW weight concentration."),
        39: ("Main 7-9; Supp 10-11", "Alternative MR estimators, leave-one-out, exposure-scale sensitivity and variance-model sensitivity."),
        40: ("Main 7-9; Supp 10-11", "Colocalization, power, minimum detectable effects and prior sensitivity."),
        41: ("Main 7; Supp 10", "Steiger direction checks are reported for all instruments; bidirectional MR was not performed."),
        42: ("N/A", "No non-MR causal estimate with a directly comparable systemic whole-blood exposure was available."),
        43: ("Main 17; Supp 10", "Figure 5 includes per-SNP and sensitivity estimates; leave-one-out ranges are tabulated."),
        45: ("Main 7, 9", "Results and Discussion summarize the null systemic blood-expression MR result against the stated objective."),
        46: ("Main 8-9", "Limitations address IV assumptions, dominant-instrument weight, exposure-scale reconstruction, tissue mismatch, ancestry and imprecision."),
        48: ("Main 8-9", "Interpretation is cautious, conditional and compared with transcriptomic and mechanistic evidence."),
        49: ("Main 8-9", "Potential mechanisms are discussed while distinguishing transcript expression, channel activity and local tissue biology; causal language is qualified."),
        50: ("Main 7-9", "Clinical interpretation is bounded to the whole-blood cis-eQTL instrument model; estimates did not support a moderate-to-large systemic effect and do not evaluate local mechanisms."),
        51: ("Main 9", "Whole-blood tissue and primarily European ancestry limit tissue, ancestry, timing and exposure-level generalizability."),
        53: ("Main 10", "No study-specific funding; source-resource funding is reported by the original studies/consortia."),
        54: ("Main 9-10; Supp 13-14", "Public source accessions, derived-result map, scripts and software versions are reported; public repository/DOI remains pending."),
        55: ("Main 10", "The author declares no competing interests."),
    }
    for row_idx, (pages, evidence) in entries.items():
        set_cell(table.rows[row_idx].cells[3], pages)
        set_cell(table.rows[row_idx].cells[4], evidence)
    clean_properties(doc, "Completed STROBE-MR checklist - PIEZO1 OA (V2)",
                     "Completed STROBE-MR reporting checklist for the corrected-SOP final-audit manuscript")
    return doc


def main() -> None:
    PACKAGE.mkdir(parents=True, exist_ok=True)
    main_doc = revise_main()
    supp_doc = revise_supplement()
    strobe_doc = build_strobe_mr_checklist()
    main_doc.save(MAIN_OUT)
    supp_doc.save(SUPP_OUT)
    strobe_doc.save(STROBE_OUT)
    print(MAIN_OUT)
    print(SUPP_OUT)
    print(STROBE_OUT)


if __name__ == "__main__":
    main()
